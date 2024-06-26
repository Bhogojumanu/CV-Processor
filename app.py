import os
import re
import streamlit as st
import xlwt
from PyPDF2 import PdfReader
import docx
import io


# Function to extract text from a DOCX file
def extract_text_from_docx(docx_file):
    doc = docx.Document(docx_file)
    text = ''
    for paragraph in doc.paragraphs:
        text += paragraph.text + '\n'
    return text


# Function to extract text from a PDF file
def extract_text_from_pdf(pdf_file):
    pdf_contents = pdf_file.read()
    with io.BytesIO(pdf_contents) as file:
        reader = PdfReader(file)
        text = ''
        for page_num in range(len(reader.pages)):
            text += reader.pages[page_num].extract_text()
        return text


# Function to extract contact information from text
def extract_contact_info(text):
    # Regular expression to match email addresses
    email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    emails = re.findall(email_pattern, text)

    # Regular expression to match phone numbers (assuming 10 digits with optional separators)
    phone_pattern = r'(?:(?:\+?(\d{1,3}))?[-.●\s]?)(?:[(]?(\d{3})[)]?[-.●\s]?)(\d{3})[-.●\s]?(\d{4})'
    phones = re.findall(phone_pattern, text)

    return emails, phones


# Function to process uploaded files and extract information
def process_uploaded_files(uploaded_files):
    if not uploaded_files:
        st.warning("No files uploaded.")
        return

    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("CV Information")
    sheet.write(0, 0, "Filename")
    sheet.write(0, 1, "Email")
    sheet.write(0, 2, "Contact Number")
    sheet.write(0, 3, "Text")
    row = 1

    for uploaded_file in uploaded_files:
        file_extension = os.path.splitext(uploaded_file.name)[1]
        if file_extension == '.docx':
            text = extract_text_from_docx(uploaded_file)
        elif file_extension == '.pdf':
            text = extract_text_from_pdf(uploaded_file)
        else:
            st.warning(f"Unsupported file format: {file_extension}")
            continue

        emails, phones = extract_contact_info(text)

        sheet.write(row, 0, uploaded_file.name)
        sheet.write(row, 1, ", ".join(emails))
        sheet.write(row, 2, ", ".join([f"{p[0]}-{p[1]}-{p[2]}-{p[3]}" for p in phones]))
        sheet.write(row, 3, text)
        row += 1

    excel_file_path = "CV_Information.xls"
    workbook.save(excel_file_path)
    st.success("Excel file created successfully.")

    # Add a download button for the Excel file
    st.markdown("### Download Excel File")
    with open(excel_file_path, "rb") as file:
        st.download_button(label="Download", data=file, file_name="CV_Information.xls", mime="application/vnd.ms-excel")


# Streamlit app
def main():
    st.title("CV Processor")

    # File uploader for user data upload
    uploaded_files = st.file_uploader("Upload your CV in DOCX or PDF file formats", type=["docx", "pdf"], accept_multiple_files=True)

    if st.button("Process"):
        process_uploaded_files(uploaded_files)


if __name__ == "__main__":
    main()
